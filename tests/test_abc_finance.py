"""ABC finance basics (A6): ledger with board approval above the limit, member accounts and receipts, voiding,
summary, Excel export, member statement, access by role."""

import io

import pytest
from openpyxl import load_workbook

from tests.conftest import add_user


@pytest.fixture
def world(platform, client_as):
    org = platform.execute("INSERT INTO organisations (name, type) VALUES ('Jiru ABC', 'abc')").lastrowid
    users = {"gm": add_user(platform, 2, phone="0911000002"), "acct": add_user(platform, 3, phone="0911000003"),
             "chair": add_user(platform, 4, phone="0911000004"), "auditor": add_user(platform, 5, phone="0911000005"),
             "store": add_user(platform, 6, phone="0911000006")}
    for uid, role in ((2, "general_manager"), (3, "accountant"), (4, "chair"), (5, "auditor"), (6, "store_keeper")):
        platform.execute("INSERT INTO abc_roles (org_id, user_id, role, since) VALUES (?, ?, ?, 'x')", (org, uid, role))
    users["member"] = add_user(platform, 10, phone="0911000010", name="Member Ten")
    platform.execute("INSERT INTO abc_members (org_id, farmer_id, status, invited_at) VALUES (?, 10, 'active', 'x')", (org,))
    platform.execute("INSERT OR IGNORE INTO abc_profiles (org_id, board_approval_limit_etb) VALUES (?, 50000)", (org,))
    platform.execute("UPDATE abc_profiles SET board_approval_limit_etb=50000 WHERE org_id=?", (org,))
    platform.commit()
    return {"org": org, "platform": platform, **{k: client_as(v) for k, v in users.items()}}


def test_expense_above_the_limit_waits_for_the_board(world):
    acct, chair = world["acct"], world["chair"]
    small = acct.post("/api/abc/finance/ledger", json={"kind": "expense", "category": "fuel", "amount_etb": 3000,
                                                      "counterparty": "Kemise fuel station", "method": "cash"}).json()
    assert small["status"] == "posted"
    big = acct.post("/api/abc/finance/ledger", json={"kind": "asset", "category": "machinery", "amount_etb": 450000,
                                                    "description": "Second-hand tractor"}).json()
    assert big["status"] == "pending" and big["approval_id"]
    s = acct.get("/api/abc/finance/summary").json()
    assert s["expense_etb"] == 3000 and s["assets_bought_etb"] == 0 and s["pending_approval_etb"] == 450000
    assert acct.post(f"/api/abc/approvals/{big['approval_id']}/approve", json={}).status_code == 403   # accountant
    assert chair.post(f"/api/abc/approvals/{big['approval_id']}/approve", json={"note": "Board 20 Oct"}).status_code == 200
    s = acct.get("/api/abc/finance/summary").json()
    assert s["assets_bought_etb"] == 450000 and s["pending_approval_etb"] == 0
    acct.post("/api/abc/finance/ledger", json={"kind": "liability", "category": "loan", "amount_etb": 200000,
                                               "direction": "in", "counterparty": "Cooperative Bank of Oromia"})
    acct.post("/api/abc/finance/ledger", json={"kind": "income", "category": "ploughing services", "amount_etb": 12000})
    s = acct.get("/api/abc/finance/summary").json()
    assert s["loans_received_etb"] == 200000 and s["result_etb"] == 9000
    assert s["cash_estimate_etb"] == 12000 - 3000 - 450000 + 200000


def test_member_account_balance_receipts_and_void(world, platform):
    acct = world["acct"]
    acct.post("/api/abc/finance/members/10", json={"kind": "input_credit", "amount_etb": 2400, "description": "Tilt 250 EC, 1 L"})
    acct.post("/api/abc/finance/members/10", json={"kind": "service_charge", "amount_etb": 1500, "description": "Ploughing 0.5 ha"})
    assert acct.post("/api/abc/finance/members/10", json={"kind": "repayment", "amount_etb": 1000,
                                                          "method": "telebirr"}).status_code == 400   # reference
    r = acct.post("/api/abc/finance/members/10", json={"kind": "repayment", "amount_etb": 1000, "method": "telebirr",
                                                       "reference": "TB998"}).json()
    assert r["balance_etb"] == 2900 and r["receipt_no"].startswith("M")
    assert acct.post("/api/abc/finance/members/10", json={"kind": "adjustment", "amount_etb": -100}).status_code == 400  # reason
    acct.post("/api/abc/finance/members/10", json={"kind": "adjustment", "amount_etb": -100, "description": "Rounding agreed"})
    st = acct.get("/api/abc/finance/members/10").json()
    assert [t["running_balance_etb"] for t in st["transactions"]] == [2400, 3900, 2900, 2800]
    accounts = acct.get("/api/abc/finance/members").json()
    assert accounts[0]["balance_etb"] == 2800 and accounts[0]["repaid_etb"] == 1000
    assert platform.execute("SELECT message_key FROM alerts WHERE farmer_id=10").fetchone()[0] == "abc_member_receipt"
    s = acct.get("/api/abc/finance/summary").json()
    assert s["income_etb"] == 1000 and s["members_owe_etb"] == 2800            # repayment booked as income
    ledger = acct.get("/api/abc/finance/ledger").json()
    linked = next(e for e in ledger if e["source"] == "member_account")
    assert acct.post(f"/api/abc/finance/ledger/{linked['id']}/void", json={"reason": "wrong"}).status_code == 400
    repayment = next(t for t in st["transactions"] if t["kind"] == "repayment")
    acct.post(f"/api/abc/finance/member-txns/{repayment['id']}/void", json={"reason": "Transfer reversed by bank"})
    s = acct.get("/api/abc/finance/summary").json()
    assert s["income_etb"] == 0 and s["members_owe_etb"] == 3800
    mine = world["member"].get("/api/abc/my/statement", params={"org_id": world["org"]}).json()
    assert mine["balance_etb"] == 3800 and len(mine["transactions"]) == 4


def test_export_and_roles(world):
    acct = world["acct"]
    acct.post("/api/abc/finance/ledger", json={"kind": "income", "category": "sales", "amount_etb": 5000})
    r = world["auditor"].get("/api/abc/finance/export.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Ledger", "Member accounts", "Member transactions", "Share payments"]
    assert wb["Ledger"]["E2"].value == 5000
    assert world["auditor"].post("/api/abc/finance/ledger", json={"kind": "income", "category": "x1", "amount_etb": 1}).status_code == 403
    assert world["store"].get("/api/abc/finance/summary").status_code == 403    # store keeper: no finance
    assert world["member"].get("/api/abc/my/statement", params={"org_id": 999}).status_code == 403
